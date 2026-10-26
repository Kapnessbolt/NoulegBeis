#!/usr/bin/env python3
"""add_video.py · skill blueage-animazioni

Incorpora un MP4 in una slide di un deck .pptx GIÀ ESISTENTE: poster, avvio automatico, loop, senza audio.
Per i deck generati dallo storyboard si usa invece il campo video di media.json della skill blueage-deck
(vedi references/powerpoint.md). XML di timing provato in PowerPoint il 14/10/2026 (deck M1 Meccanotecnica).

Uso:
  python3 add_video.py IN.pptx OUT.pptx --slide 31 --video anim1.mp4 --poster anim1_poster.png \
      --x 5.05 --y 1.4 --w 4.45 [--togli immagine,Shape16] [--link URL --testo-link "▶  Il prodotto reale"]

  --slide        numero della slide (da 1)
  --x --y --w    posizione e larghezza in pollici (altezza = w * 9/16)
  --togli        nomi delle forme da eliminare (es. l'immagine che il video sostituisce), separati da virgola
  --link         facoltativo: link al video del prodotto reale, in una riga sotto il video
OUT deve essere diverso da IN: una versione consegnata non si sovrascrive.
"""
import argparse, json, os, subprocess, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from lxml import etree

P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
TIMING = '''<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="5" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="6" dur="DUR" fill="hold"/><p:tgtEl><p:spTgt spid="SPID"/></p:tgtEl></p:cBhvr></p:cmd>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>
</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>
<p:video><p:cMediaNode vol="0" mute="1" showWhenStopped="1"><p:cTn id="7" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="SPID"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''


def durata_ms(mp4):
    out = subprocess.run(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'json', mp4],
                         capture_output=True, text=True, check=True).stdout
    return int(round(float(json.loads(out)['format']['duration']) * 1000))


def metti_video(slide, mp4, poster, x, y, w):
    h = w * 9 / 16
    mv = slide.shapes.add_movie(mp4, Inches(x), Inches(y), Inches(w), Inches(h),
                                poster_frame_image=poster, mime_type='video/mp4')
    mv.name = 'video'
    sld = slide._element
    old = sld.find(P + 'timing')
    if old is not None:
        sld.remove(old)
    t = etree.fromstring(TIMING.replace('SPID', str(mv.shape_id)).replace('DUR', str(durata_ms(mp4))))
    ext = sld.find(P + 'extLst')  # p:timing va prima di extLst
    if ext is not None:
        ext.addprevious(t)
    else:
        sld.append(t)
    return mv, h


def metti_link(slide, url, testo, x, y, w):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.3))
    tb.name = 'link-video'
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    r = tf.paragraphs[0].add_run()
    r.text = testo
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.name = 'Plus Jakarta Sans'
    r.font.color.rgb = RGBColor(0x2F, 0x80, 0xED)
    r.hyperlink.address = url


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('inp'); ap.add_argument('out')
    ap.add_argument('--slide', type=int, required=True)
    ap.add_argument('--video', required=True); ap.add_argument('--poster', required=True)
    ap.add_argument('--x', type=float, required=True); ap.add_argument('--y', type=float, required=True)
    ap.add_argument('--w', type=float, required=True)
    ap.add_argument('--togli', default='')
    ap.add_argument('--link', default=None)
    ap.add_argument('--testo-link', default='▶  Il prodotto reale')
    a = ap.parse_args()
    if os.path.abspath(a.inp) == os.path.abspath(a.out):
        sys.exit('OUT deve essere diverso da IN: una versione consegnata non si sovrascrive')
    p = Presentation(a.inp)
    s = p.slides[a.slide - 1]
    nomi = [n.strip() for n in a.togli.split(',') if n.strip()]
    tolte = 0
    for sh in list(s.shapes):
        if sh.name in nomi:
            sh._element.getparent().remove(sh._element); tolte += 1
    if nomi and tolte == 0:
        print('Attenzione: nessuna forma con i nomi', nomi, '· forme presenti:', [sh.name for sh in s.shapes])
    _, h = metti_video(s, a.video, a.poster, a.x, a.y, a.w)
    if a.link:
        metti_link(s, a.link, a.testo_link, a.x, a.y + h + 0.08, a.w)
    p.save(a.out)
    print(f'ok {a.out} · slide {a.slide}: video {os.path.basename(a.video)} ({durata_ms(a.video)/1000:.1f} s), forme tolte {tolte}'
          + (', link al prodotto reale' if a.link else ''))


if __name__ == '__main__':
    main()
